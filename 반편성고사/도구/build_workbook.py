import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from problems import M1, M2, PTS, points
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.comments import Comment

OUT = sys.argv[1]
FONT = "맑은 고딕"
N_Q, N_MC = 20, 14            # 20문항: 1~14 객관식, 15~20 단답형
FIRST, LAST = 7, 86           # 학생 데이터 행 (80명 수용)
QC = [L(4 + i) for i in range(N_Q)]   # D..W
QF, QL = QC[0], QC[-1]
CORRECT, TOTAL = "X", "Y"
RNG = f"${FIRST}:${LAST}"

def f(bold=False, size=10, color="000000"):
    return Font(name=FONT, bold=bold, size=size, color=color)

INPUT = PatternFill("solid", fgColor="FFF2CC")
HEAD = PatternFill("solid", fgColor="1F3864")
SUB = PatternFill("solid", fgColor="D9E1F2")
GREY = PatternFill("solid", fgColor="F2F2F2")
CLS = ["DDEBF7", "E2EFDA", "FCE4D6"]
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
C = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

def cell(ws, ref, v=None, font=None, fill=None, align=C, border=True, fmt=None):
    c = ws[ref]
    if v is not None:
        c.value = v
    c.font = font or f()
    if fill:
        c.fill = fill
    c.alignment = align
    if border:
        c.border = BOX
    if fmt:
        c.number_format = fmt
    return c

def title(ws, text, sub, span):
    ws.merge_cells(f"A1:{span}1")
    cell(ws, "A1", text, f(True, 14, "1F3864"), align=LEFT, border=False)
    ws.merge_cells(f"A2:{span}2")
    cell(ws, "A2", sub, f(size=9, color="595959"), align=LEFT, border=False)
    ws.row_dimensions[1].height = 26
    ws.row_dimensions[2].height = 30

def header(ws, row, labels, start=1):
    for i, lab in enumerate(labels):
        cell(ws, f"{L(start + i)}{row}", lab, f(True, 10, "FFFFFF"), HEAD)

wb = Workbook()

# ---------------- 설정 ----------------
s = wb.active
s.title = "설정"
title(s, "예비고1 반편성고사 — 설정",
      "노란 칸만 입력하세요. 반 이름·정원·반영비율·단원명을 바꾸면 모든 시트에 자동 반영됩니다. "
      "정답·배점·단원·난이도는 '공수1', '공수2' 시트 위쪽(2~5행)에서 입력합니다.", "F")
s.column_dimensions["A"].width = 14
s.column_dimensions["B"].width = 22
s.column_dimensions["C"].width = 12
s.column_dimensions["D"].width = 12
s.column_dimensions["E"].width = 44

cell(s, "A4", "① 반 설정", f(True, 11, "1F3864"), align=LEFT, border=False)
header(s, 5, ["반", "반 이름", "정원(명)"])
for i, (nm, cap) in enumerate([("심화반", 20), ("정규반", 20), ("기본반", 20)]):
    r = 6 + i
    cell(s, f"A{r}", f"{i + 1}반(성적순)", fill=GREY)
    cell(s, f"B{r}", nm, fill=INPUT)
    cell(s, f"C{r}", cap, fill=INPUT)
cell(s, "A9", "합계", f(True), GREY)
cell(s, "B9", "", fill=GREY)
cell(s, "C9", "=SUM(C6:C8)", f(True), GREY)
cell(s, "E6", "총점 석차 순으로 1반 정원만큼 → 1반, 다음 정원만큼 → 2반, 나머지 → 3반. "
     "동점자가 경계에 걸리면 모두 위 반으로 들어갑니다.", f(size=9, color="595959"), align=LEFT, border=False)
s.merge_cells("E6:F8")

cell(s, "A11", "② 점수 반영", f(True, 11, "1F3864"), align=LEFT, border=False)
header(s, 12, ["항목", "값"], start=2)
for r, lab, v, note in [
    (13, "공통수학1 반영비율", 1, "환산총점 = 공수1 점수×비율 + 공수2 점수×비율"),
    (14, "공통수학2 반영비율", 1, "예) 공수1을 더 중시하려면 1.5 / 1 처럼 입력"),
    (15, "경계 범위(점)", 5, "반 컷 점수와의 차이가 이 값 이내면 '경계'로 표시 → 상담·재검토 대상"),
]:
    cell(s, f"B{r}", lab, fill=GREY)
    cell(s, f"C{r}", v, fill=INPUT)
    cell(s, f"E{r}", note, f(size=9, color="595959"), align=LEFT, border=False)

cell(s, "A17", "③ 단원 목록 · 출제 설계 확인", f(True, 11, "1F3864"), align=LEFT, border=False)
header(s, 18, ["과목", "단원", "문항 수", "배점 합", "포함 내용(2022 개정)"])
UNITS1 = [("다항식", "다항식의 연산, 나머지정리, 인수분해"),
          ("방정식과 부등식", "복소수, 이차방정식, 이차함수, 여러 가지 방정식·부등식"),
          ("경우의 수", "합·곱의 법칙, 순열, 조합"),
          ("행렬", "행렬과 그 연산")]
UNITS2 = [("도형의 방정식", "평면좌표, 직선·원의 방정식, 도형의 이동"),
          ("집합과 명제", "집합, 명제"),
          ("함수와 그래프", "함수, 합성·역함수, 유리·무리함수")]
row = 19
UNIT_ROWS = {}
for subj, units, sh in [("공통수학1", UNITS1, "공수1"), ("공통수학2", UNITS2, "공수2")]:
    r0 = row
    for nm, desc in units:
        cell(s, f"A{row}", subj, fill=GREY)
        cell(s, f"B{row}", nm, fill=INPUT)
        cell(s, f"C{row}", f"=COUNTIF('{sh}'!${QF}$4:${QL}$4,B{row})")
        cell(s, f"D{row}", f"=SUMIF('{sh}'!${QF}$4:${QL}$4,B{row},'{sh}'!${QF}$3:${QL}$3)")
        cell(s, f"E{row}", desc, f(size=9), align=LEFT)
        row += 1
    UNIT_ROWS[sh] = (r0, row - 1)
    cell(s, f"A{row}", subj, f(True), GREY)
    cell(s, f"B{row}", "합계", f(True), GREY)
    cell(s, f"C{row}", f"=SUM(C{r0}:C{row - 1})", f(True), GREY)
    cell(s, f"D{row}", f"=SUM(D{r0}:D{row - 1})", f(True), GREY)
    cell(s, f"E{row}", "← 문항 수 20, 배점 합 100인지 확인", f(size=9, color="C00000"), align=LEFT)
    row += 1


# ---- 보고서 설정 (make_reports.py 가 이 칸 위치를 읽음) ----
REPORT_ROW = 30
cell(s, f"A{REPORT_ROW}", "④ 학부모 보고서 설정", f(True, 11, "1F3864"), align=LEFT, border=False)
header(s, REPORT_ROW + 1, ["항목", "값"], start=1)
s.merge_cells(f"B{REPORT_ROW + 1}:E{REPORT_ROW + 1}")
REPORT_ITEMS = [
    ("학원명", "OO수학학원"),
    ("보고서 제목", "예비고1 반편성 진단평가 결과 보고서"),
    ("시험일", "2026년 12월 00일"),
    ("전체 평균 표시", "표시"),
    ("'매우 우수' 기준(득점률 %)", 85),
    ("'우수' 기준(득점률 %)", 70),
    ("'양호' 기준(득점률 %)", 50),
    ("상담 안내 문구", "보고서를 받아보신 뒤 개별 상담을 진행합니다. 상담 일정은 담당 선생님께서 개별적으로 연락드리겠습니다."),
    ("학원 연락처", "02-000-0000"),
]
dv_show = DataValidation(type="list", formula1='"표시,숨김"', allow_blank=False)
s.add_data_validation(dv_show)
for i, (lab, v) in enumerate(REPORT_ITEMS):
    r = REPORT_ROW + 2 + i
    cell(s, f"A{r}", lab, f(size=9), GREY, align=LEFT)
    s.merge_cells(f"B{r}:E{r}")
    cell(s, f"B{r}", v, fill=INPUT, align=LEFT)
    if lab == "전체 평균 표시":
        dv_show.add(f"B{r}")
cell(s, f"F{REPORT_ROW + 5}", "'숨김'이면 보고서에 평균을 싣지 않음", f(size=9, color="595959"), align=LEFT, border=False)
cell(s, f"F{REPORT_ROW + 8}", "득점률이 '양호' 기준 미만이면 '보완 필요'", f(size=9, color="595959"), align=LEFT, border=False)
s.row_dimensions[REPORT_ROW + 9].height = 30
s.column_dimensions["F"].width = 34

CLASS_ROW = REPORT_ROW + 13
cell(s, f"A{CLASS_ROW}", "⑤ 반 소개 (보고서의 추천 반 설명)", f(True, 11, "1F3864"), align=LEFT, border=False)
header(s, CLASS_ROW + 1, ["반 이름", "소개 문구"])
s.merge_cells(f"B{CLASS_ROW + 1}:E{CLASS_ROW + 1}")
CLASS_DESC = [
    "공통수학1·2 개념을 빠르게 정리하고, 고난도 문항과 고1 내신 최상위권 대비까지 진행하는 반입니다.",
    "핵심 개념을 탄탄히 다지면서 내신 기출 유형과 중상위 난도 문항을 체계적으로 훈련하는 반입니다.",
    "중등 연계 개념부터 차근차근 점검하며, 공통수학의 핵심 개념과 기본 유형을 확실히 완성하는 반입니다.",
]
for i, d in enumerate(CLASS_DESC):
    r = CLASS_ROW + 2 + i
    cell(s, f"A{r}", f"=B{6 + i}", f(True), GREY)
    s.merge_cells(f"B{r}:E{r}")
    cell(s, f"B{r}", d, fill=INPUT, align=LEFT)
    s.row_dimensions[r].height = 30

ADV_ROW = CLASS_ROW + 6
cell(s, f"A{ADV_ROW}", "⑥ 단원별 학습 제언 (보완이 필요한 단원에 자동 기재)", f(True, 11, "1F3864"), align=LEFT, border=False)
header(s, ADV_ROW + 1, ["단원", "학습 제언 문구"])
s.merge_cells(f"B{ADV_ROW + 1}:E{ADV_ROW + 1}")
ADVICE = {
    "다항식": "다항식의 연산과 나머지정리, 인수분해 공식을 정확하고 빠르게 적용하는 연습이 필요합니다.",
    "방정식과 부등식": "복소수 계산과 이차방정식·이차함수의 관계를 그래프와 연결해 이해하는 학습이 필요합니다.",
    "경우의 수": "합·곱의 법칙과 순열·조합을 구분하는 기준을 정리하고, 빠짐없이 중복 없이 세는 연습이 필요합니다.",
    "행렬": "행렬의 연산 규칙을 정확히 익히고, 계산 과정에서의 실수를 줄이는 연습이 필요합니다.",
    "도형의 방정식": "좌표와 식을 연결하는 연습, 특히 직선·원의 방정식과 도형의 이동을 그림과 함께 해석하는 훈련이 필요합니다.",
    "집합과 명제": "용어와 기호의 정확한 의미를 정리하고, 명제의 참·거짓을 논리적으로 판단하는 연습이 필요합니다.",
    "함수와 그래프": "함수의 정의와 합성함수·역함수 개념을 정리하고, 유리·무리함수의 그래프를 직접 그려 보는 학습이 필요합니다.",
}
unit_src_rows = list(range(UNIT_ROWS["공수1"][0], UNIT_ROWS["공수1"][1] + 1)) + \
                list(range(UNIT_ROWS["공수2"][0], UNIT_ROWS["공수2"][1] + 1))
for i, ur in enumerate(unit_src_rows):
    r = ADV_ROW + 2 + i
    cell(s, f"A{r}", f"=B{ur}", f(size=9), GREY)
    s.merge_cells(f"B{r}:E{r}")
    cell(s, f"B{r}", ADVICE[s[f"B{ur}"].value], fill=INPUT, align=LEFT)
    s.row_dimensions[r].height = 30
s.column_dimensions["A"].width = 24

# ---------------- 학생명단 ----------------
st = wb.create_sheet("학생명단")
title(st, "학생 명단", "노란 칸에 학생 정보를 입력하세요. 여기 입력한 순서(행)가 공수1·공수2·반배정 시트의 행과 "
      "그대로 연결됩니다. 7행은 예시이니 지우고 사용하세요. G열 '선생님 한마디'는 학부모 보고서에 그대로 실립니다(비워두면 생략).", "G")
header(st, 6, ["No", "이름", "학교", "구분", "연락처", "비고(내부용)", "선생님 한마디(보고서용)"])
for col, w in zip("ABCDEFG", [6, 14, 14, 8, 16, 20, 60]):
    st.column_dimensions[col].width = w
dv_kind = DataValidation(type="list", formula1='"재원,외부"', allow_blank=True)
st.add_data_validation(dv_kind)
for r in range(FIRST, LAST + 1):
    cell(st, f"A{r}", r - FIRST + 1, fill=GREY)
    for col in "BCDEFG":
        cell(st, f"{col}{r}", fill=INPUT, align=LEFT if col in "FG" else C)
dv_kind.add(f"D{FIRST}:D{LAST}")
for col, v in zip("BCDEFG", ["홍길동(예시)", "OO중", "외부", "010-0000-0000", "예시 행 — 삭제 후 사용",
                              "수업 태도가 성실하고 끝까지 풀어내려는 끈기가 돋보입니다."]):
    st[f"{col}{FIRST}"].value = v
st.freeze_panes = f"C{FIRST}"

# ---------------- 채점 시트 ----------------
# 시험지(problems.py)의 정답·배점·단원·난이도를 기본값으로 사용
EXAM = {"공수1": M1, "공수2": M2}

def score_sheet(name, subj):
    M = EXAM[name]
    ws = wb.create_sheet(name)
    u0, u1 = UNIT_ROWS[name]
    unit_cols = [L(26 + i) for i in range(u1 - u0 + 1)]
    last_col = unit_cols[-1]
    ws.merge_cells("A1:C1")
    cell(ws, "A1", f"{subj} 채점표", f(True, 14, "1F3864"), align=LEFT, border=False)
    ws.merge_cells(f"D1:{last_col}1")
    cell(ws, "D1", "① 2~5행(노란 칸)에 정답·배점·단원·난이도 입력  ② 학생 행에 학생이 쓴 답을 그대로 입력"
         "(객관식 1~5, 단답형 0~999) → 틀린 답은 빨간 칸, 점수 자동 계산. 결시생은 비워두면 미응시 처리.",
         f(size=9, color="595959"), align=LEFT, border=False)
    ws.row_dimensions[1].height = 34
    ws.merge_cells("A2:C2")
    ws.merge_cells("A3:C3")
    ws.merge_cells("A4:C4")
    ws.merge_cells("A5:C5")
    for r, lab in [(2, "정답 ▶"), (3, "배점 ▶"), (4, "단원 ▶"), (5, "난이도 ▶")]:
        cell(ws, f"A{r}", lab, f(True), SUB)
    dv_unit = DataValidation(type="list", formula1=f"='설정'!$B${u0}:$B${u1}", allow_blank=True)
    dv_lv = DataValidation(type="list", formula1='"하,중하,중,중상,상,최상"', allow_blank=True)
    dv_mc = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True,
                           error="객관식은 1~5 사이 정수만 입력", errorTitle="입력 오류", showErrorMessage=True)
    dv_sa = DataValidation(type="whole", operator="between", formula1="0", formula2="999", allow_blank=True,
                           error="단답형은 0~999 사이 정수만 입력", errorTitle="입력 오류", showErrorMessage=True)
    for dv in (dv_unit, dv_lv, dv_mc, dv_sa):
        ws.add_data_validation(dv)
    for i, c in enumerate(QC):
        mc = i < N_MC
        cell(ws, f"{c}2", M[i]["ans"], f(True, 10, "C00000"), INPUT)
        cell(ws, f"{c}3", points(M[i]), fill=INPUT)
        cell(ws, f"{c}4", M[i]["unit"], f(size=8), INPUT)
        cell(ws, f"{c}5", M[i]["level"], fill=INPUT)
        cell(ws, f"{c}6", f"{i + 1}\n{'객관식' if mc else '단답형'}", f(True, 9, "FFFFFF"),
             HEAD if mc else PatternFill("solid", fgColor="7030A0"))
        ws.column_dimensions[c].width = 6.2
        (dv_mc if mc else dv_sa).add(f"{c}2")
        (dv_mc if mc else dv_sa).add(f"{c}{FIRST}:{c}{LAST}")
    dv_unit.add(f"{QF}4:{QL}4")
    dv_lv.add(f"{QF}5:{QL}5")
    ws.row_dimensions[4].height = 42
    ws.row_dimensions[6].height = 30
    header(ws, 6, ["No", "이름", "학교"])
    header(ws, 6, ["맞힌\n개수", "총점"], start=24)
    for i, uc in enumerate(unit_cols):
        cell(ws, f"{uc}6", f"='설정'!B{u0 + i}", f(True, 9, "FFFFFF"), PatternFill("solid", fgColor="2F5597"))
        cell(ws, f"{uc}3", f"=SUMIF(${QF}$4:${QL}$4,{uc}$6,${QF}$3:${QL}$3)", f(size=9), GREY)
        cell(ws, f"{uc}5", "만점 ▲", f(size=8, color="595959"), border=False)
        ws.column_dimensions[uc].width = 10
    cell(ws, "X3", f"=COUNT({QF}2:{QL}2)", f(size=9), GREY)
    cell(ws, "Y3", f"=SUM({QF}3:{QL}3)", f(True, 9), GREY)
    cell(ws, "X5", "정답입력수▲", f(size=8, color="595959"), border=False)
    cell(ws, "Y5", "만점 ▲", f(size=8, color="595959"), border=False)
    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["X"].width = 7
    ws.column_dimensions["Y"].width = 8
    key = f"${QF}$2:${QL}$2"
    pts = f"${QF}$3:${QL}$3"
    for r in range(FIRST, LAST + 1):
        resp = f"{QF}{r}:{QL}{r}"
        cell(ws, f"A{r}", r - FIRST + 1, fill=GREY)
        cell(ws, f"B{r}", f"=IF('학생명단'!B{r}=\"\",\"\",'학생명단'!B{r})", fill=GREY)
        cell(ws, f"C{r}", f"=IF('학생명단'!C{r}=\"\",\"\",'학생명단'!C{r})", f(size=9), GREY)
        for c in QC:
            cell(ws, f"{c}{r}", fill=INPUT)
        cell(ws, f"X{r}", f"=IF(COUNTA({resp})=0,\"\",SUMPRODUCT(({resp}={key})*({resp}<>\"\")))")
        cell(ws, f"Y{r}", f"=IF(COUNTA({resp})=0,\"\",SUMPRODUCT(({resp}={key})*({resp}<>\"\")*{pts}))", f(True))
        for uc in unit_cols:
            cell(ws, f"{uc}{r}",
                 f"=IF($Y{r}=\"\",\"\",SUMPRODUCT((${QF}{r}:${QL}{r}={key})*(${QF}{r}:${QL}{r}<>\"\")"
                 f"*{pts}*(${QF}$4:${QL}$4={uc}$6)))")
    ws.conditional_formatting.add(
        f"{QF}{FIRST}:{QL}{LAST}",
        FormulaRule(formula=[f"AND({QF}{FIRST}<>\"\",{QF}$2<>\"\",{QF}{FIRST}<>{QF}$2)"],
                    fill=PatternFill("solid", fgColor="F8CBAD"), font=Font(color="C00000", bold=True)))
    ws.conditional_formatting.add(
        f"{QF}2:{QL}2",
        FormulaRule(formula=[f"{QF}2=\"\""], fill=PatternFill("solid", fgColor="FF9999")))
    ws.freeze_panes = f"D{FIRST}"
    return ws

score_sheet("공수1", "공통수학1")
score_sheet("공수2", "공통수학2")

# ---------------- 반배정 ----------------
b = wb.create_sheet("반배정")
title(b, "반 배정", "자동 계산 시트입니다. 'K열 수동조정'에서만 반을 바꿀 수 있고(상담 결과 반영 등), "
      "최종반 = 수동조정이 있으면 그 값, 없으면 자동배정. 주황색 행 = 반 경계 점수대 학생.", "U")
header(b, 6, ["No", "이름", "학교", "구분", "공수1", "공수2", "환산총점", "석차", "자동배정",
              "경계", "수동조정", "최종반", "비고", "정렬키"])
for col, w in zip("ABCDEFGHIJKLMN", [5, 12, 11, 7, 7, 7, 9, 6, 10, 6, 10, 10, 13, 8]):
    b.column_dimensions[col].width = w
b.column_dimensions["N"].hidden = True
NAMES = ["'설정'!$B$6", "'설정'!$B$7", "'설정'!$B$8"]
CAP1, CAP2 = "'설정'!$C$6", "'설정'!$C$7"
G = f"$G${FIRST}:$G${LAST}"
I = f"$I${FIRST}:$I${LAST}"
Lr = f"$L${FIRST}:$L${LAST}"
D = f"$D${FIRST}:$D${LAST}"
CUT1, CUT2 = "$T$7", "$T$8"
dv_cls = DataValidation(type="list", formula1="='설정'!$B$6:$B$8", allow_blank=True)
b.add_data_validation(dv_cls)
for r in range(FIRST, LAST + 1):
    cell(b, f"A{r}", r - FIRST + 1, fill=GREY)
    for col, src in zip("BCD", "BCD"):
        cell(b, f"{col}{r}", f"=IF('학생명단'!{src}{r}=\"\",\"\",'학생명단'!{src}{r})")
    cell(b, f"E{r}", f"='공수1'!Y{r}")
    cell(b, f"F{r}", f"='공수2'!Y{r}")
    cell(b, f"G{r}", f"=IF(AND(E{r}=\"\",F{r}=\"\"),\"\",N(E{r})*'설정'!$C$13+N(F{r})*'설정'!$C$14)",
         f(True), fmt="0.#")
    cell(b, f"H{r}", f"=IF(G{r}=\"\",\"\",COUNTIF({G},\">\"&G{r})+1)")
    cell(b, f"I{r}", f"=IF(H{r}=\"\",\"\",IF(H{r}<={CAP1},{NAMES[0]},IF(H{r}<={CAP1}+{CAP2},{NAMES[1]},{NAMES[2]})))")
    cell(b, f"J{r}", f"=IF(G{r}=\"\",\"\",IF(OR(AND(ISNUMBER({CUT1}),ABS(G{r}-{CUT1})<='설정'!$C$15),"
                     f"AND(ISNUMBER({CUT2}),ABS(G{r}-{CUT2})<='설정'!$C$15)),\"경계\",\"\"))", f(True, 10, "C55A11"))
    cell(b, f"K{r}", fill=INPUT)
    cell(b, f"L{r}", f"=IF(K{r}<>\"\",K{r},I{r})", f(True))
    cell(b, f"M{r}", f"=IF(G{r}=\"\",\"\",IF(E{r}=\"\",\"공수1 미응시\",IF(F{r}=\"\",\"공수2 미응시\",\"\")))",
         f(size=9, color="C00000"))
    cell(b, f"N{r}", f"=IF(H{r}=\"\",\"\",H{r}+ROW()/100000)")
dv_cls.add(f"K{FIRST}:K{LAST}")
b.conditional_formatting.add(f"A{FIRST}:M{LAST}",
                             FormulaRule(formula=[f"$J{FIRST}=\"경계\""], fill=PatternFill("solid", fgColor="FBE5D6")))
for i, nm in enumerate(NAMES):
    b.conditional_formatting.add(f"L{FIRST}:L{LAST}",
                                 FormulaRule(formula=[f"AND($L{FIRST}<>\"\",$L{FIRST}={nm})"],
                                             fill=PatternFill("solid", fgColor=CLS[i])))
# 요약표
cell(b, "P5", "반별 요약", f(True, 11, "1F3864"), align=LEFT, border=False)
header(b, 6, ["반", "정원", "자동\n인원", "최종\n인원", "자동컷\n(최저점)", "최종\n평균", "재원", "외부"], start=16)
for i, nm in enumerate(NAMES):
    r = 7 + i
    cell(b, f"P{r}", f"={nm}", f(True), PatternFill("solid", fgColor=CLS[i]))
    cell(b, f"Q{r}", f"='설정'!$C${6 + i}")
    cell(b, f"R{r}", f"=COUNTIF({I},P{r})")
    cell(b, f"S{r}", f"=COUNTIF({Lr},P{r})", f(True))
    cell(b, f"T{r}", f"=IF(R{r}=0,\"\",_xlfn.MINIFS({G},{I},P{r}))", fmt="0.#")
    cell(b, f"U{r}", f"=IFERROR(AVERAGEIF({Lr},P{r},{G}),\"\")", fmt="0.0")
    cell(b, f"V{r}", f"=COUNTIFS({Lr},P{r},{D},\"재원\")")
    cell(b, f"W{r}", f"=COUNTIFS({Lr},P{r},{D},\"외부\")")
cell(b, "P10", "합계", f(True), GREY)
for col in "QRSVW":
    cell(b, f"{col}10", f"=SUM({col}7:{col}9)", f(True), GREY)
cell(b, "T10", "", fill=GREY)
cell(b, "U10", f"=IFERROR(AVERAGE({G}),\"\")", f(True), GREY, fmt="0.0")
cell(b, "P12", "· 경계 = 1·2반 자동컷 점수와 차이가 '설정' 경계 범위 이내인 학생\n"
     "· 수동조정한 학생은 최종 인원·평균에 반영됩니다(자동컷은 그대로)", f(size=9, color="595959"),
     align=LEFT, border=False)
b.merge_cells("P12:W14")
for col in "PQRSTUVW":
    b.column_dimensions[col].width = 8
b.column_dimensions["P"].width = 10
b.freeze_panes = f"C{FIRST}"

# ---------------- 석차순 ----------------
k = wb.create_sheet("석차순")
title(k, "석차순 명단 (발표·상담용)", "자동 정렬됩니다. 인쇄 시 그대로 사용하세요.", "J")
header(k, 6, ["석차", "이름", "학교", "구분", "공수1", "공수2", "환산총점", "최종반", "경계", "비고"])
for col, w in zip("ABCDEFGHIJ", [6, 12, 12, 7, 7, 7, 9, 10, 6, 13]):
    k.column_dimensions[col].width = w
k.column_dimensions["K"].hidden = True
k.column_dimensions["L"].hidden = True
KEYS = f"'반배정'!$N${FIRST}:$N${LAST}"
for r in range(FIRST, LAST + 1):
    cell(k, f"K{r}", f"=IFERROR(SMALL({KEYS},{r - FIRST + 1}),\"\")")
    cell(k, f"L{r}", f"=IF(K{r}=\"\",\"\",MATCH(K{r},{KEYS},0))")
    for col, src in zip("ABCDEFGHIJ", "HBCDEFGLJM"):
        cell(k, f"{col}{r}", f"=IF($L{r}=\"\",\"\",INDEX('반배정'!{src}${FIRST}:{src}${LAST},$L{r}))",
             f(True) if col in "GH" else None, fmt="0.#" if col == "G" else None)
for i, nm in enumerate(NAMES):
    k.conditional_formatting.add(f"A{FIRST}:J{LAST}",
                                 FormulaRule(formula=[f"AND($H{FIRST}<>\"\",$H{FIRST}={nm})"],
                                             fill=PatternFill("solid", fgColor=CLS[i])))
k.freeze_panes = f"A{FIRST}"

# ---------------- 문항분석 ----------------
a = wb.create_sheet("문항분석")
title(a, "문항 분석", "변별도 = 1반 정답률 − 3반 정답률 (0.3 이상이면 반을 잘 가르는 문항). "
      "객관식은 선지별 선택 인원으로 매력적 오답을 확인하세요.", "R")
cols = ["문항", "유형", "단원", "난이도", "배점", "정답", "응시", "정답자", "정답률",
        "1반\n정답률", "2반\n정답률", "3반\n정답률", "변별도", "진단", "①", "②", "③", "④", "⑤"]
widths = [5, 7, 14, 6, 5, 5, 5, 6, 7, 8, 8, 8, 7, 14, 5, 5, 5, 5, 5]
for i, w in enumerate(widths):
    a.column_dimensions[L(i + 1)].width = w
row = 4
for sh, subj in [("공수1", "공통수학1"), ("공수2", "공통수학2")]:
    a.merge_cells(f"A{row}:E{row}")
    cell(a, f"A{row}", subj, f(True, 11, "1F3864"), align=Alignment(horizontal="left", vertical="center"), border=False)
    row += 1
    header(a, row, cols)
    for j, nm in enumerate(NAMES):
        a[f"{L(10 + j)}{row}"].value = f"={nm}&\" 정답률\""
    a.row_dimensions[row].height = 30
    row += 1
    for q in range(N_Q):
        c = QC[q]
        col = f"'{sh}'!${c}${FIRST}:${c}${LAST}"
        key = f"'{sh}'!${c}$2"
        tot = f"'{sh}'!$Y${FIRST}:$Y${LAST}"
        r = row
        cell(a, f"A{r}", q + 1, f(True), GREY)
        cell(a, f"B{r}", "객관식" if q < N_MC else "단답형", f(size=9), GREY)
        cell(a, f"C{r}", f"='{sh}'!{c}4", f(size=9))
        cell(a, f"D{r}", f"='{sh}'!{c}5")
        cell(a, f"E{r}", f"='{sh}'!{c}3")
        cell(a, f"F{r}", f"=IF({key}=\"\",\"\",{key})", f(True, 10, "C00000"))
        cell(a, f"G{r}", f"=COUNT({tot})")
        cell(a, f"H{r}", f"=IF(F{r}=\"\",\"\",SUMPRODUCT(({col}={key})*({col}<>\"\")))")
        cell(a, f"I{r}", f"=IF(OR(F{r}=\"\",G{r}=0),\"\",H{r}/G{r})", f(True), fmt="0%")
        for j, nm in enumerate(NAMES):
            cls = f"'반배정'!$I${FIRST}:$I${LAST}"
            cell(a, f"{L(10 + j)}{r}",
                 f"=IF(F{r}=\"\",\"\",IFERROR(SUMPRODUCT(({col}={key})*({col}<>\"\")*({cls}={nm}))"
                 f"/COUNTIFS({cls},{nm},{tot},\">=0\"),\"\"))", fmt="0%")
        cell(a, f"M{r}", f"=IF(OR(J{r}=\"\",L{r}=\"\"),\"\",J{r}-L{r})", f(True), fmt="0.00")
        cell(a, f"N{r}", f"=IF(I{r}=\"\",\"\",IF(AND(M{r}<>\"\",M{r}<0),\"역변별·문항검토\","
                         f"IF(I{r}>=0.9,\"매우 쉬움\",IF(I{r}<=0.15,\"매우 어려움\","
                         f"IF(AND(M{r}<>\"\",M{r}<0.2),\"변별력 낮음\",\"양호\")))))", f(size=9))
        for o in range(5):
            if q < N_MC:
                cell(a, f"{L(15 + o)}{r}", f"=IF(G{r}=0,\"\",COUNTIF({col},{o + 1}))", f(size=9))
            else:
                cell(a, f"{L(15 + o)}{r}", "-", f(size=9, color="A6A6A6"), GREY)
        row += 1
    a.conditional_formatting.add(f"N{row - N_Q}:N{row - 1}",
        FormulaRule(formula=[f"N{row - N_Q}=\"양호\""], fill=PatternFill("solid", fgColor="E2EFDA")))
    a.conditional_formatting.add(f"N{row - N_Q}:N{row - 1}",
        FormulaRule(formula=[f"AND(N{row - N_Q}<>\"\",N{row - N_Q}<>\"양호\")"],
                    fill=PatternFill("solid", fgColor="FCE4D6")))
    # 정답 선지 강조
    a.conditional_formatting.add(f"O{row - N_Q}:S{row - N_Q + N_MC - 1}",
        FormulaRule(formula=[f"COLUMN(O{row - N_Q})-14=$F{row - N_Q}"],
                    fill=PatternFill("solid", fgColor="C6E0B4"), font=Font(bold=True)))
    row += 2
a.freeze_panes = "A4"

for ws in wb.worksheets:
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

wb.save(OUT)
print("saved", OUT)
