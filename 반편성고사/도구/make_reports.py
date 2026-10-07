# -*- coding: utf-8 -*-
"""채점 엑셀(입력값)을 읽어 학부모 보고서와 상담 카드(PDF)를 만든다.

사용: python make_reports.py 채점파일.xlsx 출력폴더
  - 학부모보고서_전체.pdf, 학부모보고서/번호_이름.pdf : 석차·백분위 없이 성취 수준, (선택)전체 평균, 추천 반(빈칸)
  - 학부모보고서_한글/전체.hml, 학부모보고서_한글/번호_이름.hml : 같은 내용의 수정 가능한 한글 파일
  - 상담카드_전체.pdf : 내부용. 석차, 반 컷과의 차이, 경계 여부, 오답 분석
엑셀의 수식 결과(캐시)에 의존하지 않고 입력값으로 직접 계산한다.
필요: openpyxl, playwright(Chromium)
"""
import sys, os, html, re
from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
try:
    from problems import M1, M2
    POINT = {("공수1", p["no"]): p["point"] for p in M1}
    POINT.update({("공수2", p["no"]): p["point"] for p in M2})
except Exception:
    POINT = {}

FIRST, LAST = 7, 86
QCOLS = list(range(4, 24))  # D..W
SUBJ = [("공수1", "공통수학1"), ("공수2", "공통수학2")]
CHROME = os.environ.get("CHROME_PATH", "/opt/pw-browsers/chromium-1194/chrome-linux/chrome")
CIRC = "①②③④⑤"


def num(v):
    if v is None or v == "":
        return None
    if isinstance(v, str):
        v = v.strip()
        try:
            f = float(v)
            return int(f) if f.is_integer() else f
        except ValueError:
            return v
    return int(v) if isinstance(v, float) and v.is_integer() else v


def read(path):
    wb = load_workbook(path)  # 입력 칸만 읽으므로 수식 캐시 불필요
    s = wb["설정"]
    val = lambda ref: s[ref].value
    cfg = {
        "classes": [val("B6"), val("B7"), val("B8"), val("B9") or "보류"],
        # 절대평가 기준: (반 이름, 공수1 이상, 공수1 미만, 공수2 이상, 공수2 미만), 빈칸은 조건 없음
        "rules": [(val("B%d" % r), num(val("C%d" % r)), num(val("D%d" % r)), num(val("E%d" % r)), num(val("F%d" % r)))
                  for r in (6, 7, 8)],
        "w": [float(val("C13") or 1), float(val("C14") or 1)],
        "border": float(val("C15") or 0),
    }
    # 보고서 설정: A열 항목명으로 찾는다
    lab = {}
    adv_row = None
    for r in range(1, s.max_row + 1):
        a = s.cell(r, 1).value
        if isinstance(a, str):
            lab.setdefault(a.strip(), s.cell(r, 2).value)
            if "학습 제언" in a and a.strip()[0] in "④⑤⑥⑦":
                adv_row = r
    g = lambda k, d="": lab.get(k) if lab.get(k) not in (None, "") else d
    cfg.update({
        "academy": g("학원명"), "title": g("보고서 제목"), "date": g("시험일"),
        "show_avg": g("전체 평균 표시", "표시") != "숨김",
        "lv": [float(g("'매우 우수' 기준(득점률 %)", 85)), float(g("'우수' 기준(득점률 %)", 70)), float(g("'양호' 기준(득점률 %)", 50))],
        "notice": g("상담 안내 문구"), "contact": g("학원 연락처"),
    })
    units = {"공수1": [], "공수2": []}
    r = 19
    for key in ("공수1", "공수2"):
        while s["B%d" % r].value != "합계":
            units[key].append(s["B%d" % r].value)
            r += 1
        r += 1
    cfg["units"] = units
    order = units["공수1"] + units["공수2"]
    cfg["advice"] = {u: (s["B%d" % (adv_row + 2 + i)].value or "") for i, u in enumerate(order)} if adv_row else {}

    exams = {}
    for key, _ in SUBJ:
        ws = wb[key]
        exams[key] = [dict(no=i + 1, key=num(ws.cell(2, c).value), pts=num(ws.cell(3, c).value) or 0,
                           unit=ws.cell(4, c).value, level=ws.cell(5, c).value, col=c) for i, c in enumerate(QCOLS)]

    st, b = wb["학생명단"], wb["반배정"]
    students = []
    for r in range(FIRST, LAST + 1):
        name = st.cell(r, 2).value
        resp = {key: [num(wb[key].cell(r, c).value) for c in QCOLS] for key, _ in SUBJ}
        if not name and not any(v is not None for key in resp for v in resp[key]):
            continue
        if name and "(예시)" in str(name):
            continue
        students.append(dict(row=r, no=r - FIRST + 1, name=str(name or "(이름 없음)"), school=st.cell(r, 3).value or "",
                             kind=st.cell(r, 4).value or "", phone=st.cell(r, 5).value or "", memo=st.cell(r, 6).value or "",
                             manual=b.cell(r, 11).value, resp=resp))
    return cfg, exams, students


def compute(cfg, exams, students):
    for s in students:
        s["took"], s["score"], s["ok"], s["unit"] = {}, {}, {}, {}
        for key, _ in SUBJ:
            resp = s["resp"][key]
            took = any(v is not None for v in resp)
            s["took"][key] = took
            ok = [took and v is not None and q["key"] is not None and v == q["key"] for v, q in zip(resp, exams[key])]
            s["ok"][key] = ok
            s["score"][key] = sum(q["pts"] for q, o in zip(exams[key], ok) if o) if took else None
            for u in cfg["units"][key]:
                full = sum(q["pts"] for q in exams[key] if q["unit"] == u)
                got = sum(q["pts"] for q, o in zip(exams[key], ok) if o and q["unit"] == u)
                s["unit"][(key, u)] = (got / full if full and took else None)
        if not any(s["took"].values()):
            s["total"] = None
        else:
            s["total"] = sum((s["score"][k] or 0) * w for (k, _), w in zip(SUBJ, cfg["w"]))
    ranked = [s for s in students if s["total"] is not None]
    hold = cfg["classes"][3]

    def fits(x, lo, hi):
        return (lo is None or x >= lo) and (hi is None or x < hi)

    for s in ranked:
        s["rank"] = 1 + sum(1 for t in ranked if t["total"] > s["total"])
        a, b = s["score"]["공수1"], s["score"]["공수2"]
        s["auto"] = hold
        if a is not None and b is not None:
            for name, lo1, hi1, lo2, hi2 in cfg["rules"]:
                if fits(a, lo1, hi1) and fits(b, lo2, hi2):
                    s["auto"] = name
                    break
        s["final"] = s["manual"] or s["auto"]
        # 경계: 과목 점수가 그 과목의 기준 점수와 경계 범위 이내
        th1 = [v for r_ in cfg["rules"] for v in r_[1:3] if v is not None]
        th2 = [v for r_ in cfg["rules"] for v in r_[3:5] if v is not None]
        s["near"] = (a is not None and b is not None and
                     (any(abs(a - v) <= cfg["border"] for v in th1) or any(abs(b - v) <= cfg["border"] for v in th2)))
    for s in students:
        s.setdefault("rank", None); s.setdefault("auto", None); s.setdefault("final", s["manual"]); s.setdefault("near", False)
    stats = {"n": len(ranked)}
    for key, _ in SUBJ:
        takers = [s for s in students if s["took"][key]]
        stats[("avg", key)] = sum(s["score"][key] for s in takers) / len(takers) if takers else None
        stats[("full", key)] = sum(q["pts"] for q in exams[key])
        for u in cfg["units"][key]:
            v = [s["unit"][(key, u)] for s in takers if s["unit"][(key, u)] is not None]
            stats[("uavg", key, u)] = sum(v) / len(v) if v else None
            for cname in cfg["classes"]:
                vv = [s["unit"][(key, u)] for s in takers if s["final"] == cname and s["unit"][(key, u)] is not None]
                stats[("cavg", key, u, cname)] = sum(vv) / len(vv) if vv else None
        for q in exams[key]:
            n = len(takers)
            q["rate"] = sum(1 for s in takers if s["ok"][key][q["no"] - 1]) / n if n else None
    return stats


def level(rate, cfg):
    if rate is None:
        return None
    p = rate * 100
    a, b, c = cfg["lv"]
    return "매우 우수" if p >= a else "우수" if p >= b else "양호" if p >= c else "보완 필요"


LVCLS = {"매우 우수": "l0", "우수": "l1", "양호": "l2", "보완 필요": "l3"}


def josa(word, pair):  # pair=("과","와") 등: 받침 있으면 앞
    ch = word.strip()[-1]
    code = ord(ch) - 0xAC00
    has = 0 <= code <= 11171 and code % 28 != 0
    return word + (pair[0] if has else pair[1])


def opinion(s, cfg, stats):
    name = s["name"]
    took = [k for k, _ in SUBJ if s["took"][k]]
    full = sum(stats[("full", k)] for k in took)
    rate = sum(s["score"][k] for k in took) / full if full else 0
    lv = level(rate, cfg)
    first = {
        "매우 우수": f"{name} 학생은 공통수학 전반에 걸쳐 개념 이해와 문제 해결력이 매우 뛰어납니다. 고난도 문항에서도 안정적으로 풀이 과정을 이어 가는 힘이 돋보였습니다.",
        "우수": f"{name} 학생은 공통수학의 핵심 개념을 잘 이해하고 있으며, 안정적인 문제 해결력을 보여 주었습니다. 일부 단원을 보완하면 한 단계 더 성장할 수 있습니다.",
        "양호": f"{name} 학생은 공통수학의 기본 개념을 갖추고 있습니다. 단원별 이해도의 차이를 줄여 나가면 실력이 더욱 단단해질 것으로 기대됩니다.",
        "보완 필요": f"{name} 학생은 이번 진단평가를 통해 앞으로 다져야 할 부분을 분명히 확인할 수 있었습니다. 핵심 개념부터 차근차근 쌓아 가면 충분히 성장할 수 있습니다.",
    }[lv]
    units = [(k, u, s["unit"][(k, u)]) for k, _ in SUBJ for u in cfg["units"][k] if s["unit"][(k, u)] is not None]
    strong = sorted([x for x in units if x[2] * 100 >= cfg["lv"][1]], key=lambda x: -x[2])[:2]
    weak = sorted([x for x in units if x[2] * 100 < cfg["lv"][1]], key=lambda x: x[2])[:2]
    parts = [first]
    if len(strong) == 2:
        parts.append(f"특히 {josa(strong[0][1], ('과', '와'))} {strong[1][1]} 단원에서 강점을 보였습니다.")
    elif len(strong) == 1:
        parts.append(f"특히 {strong[0][1]} 단원에서 강점을 보였습니다.")
    missing = [n for k, n in SUBJ if not s["took"][k]]
    if missing:
        parts.append(f"{'·'.join(missing)}는 응시하지 않아 응시한 과목을 중심으로 분석하였습니다.")
    tips = [(u, cfg["advice"].get(u, "")) for _, u, _ in weak]
    return " ".join(parts), tips


CSS = """
@page{size:A4;margin:0}
*{box-sizing:border-box}
body{margin:0;font-family:"Noto Sans KR","Malgun Gothic","Apple SD Gothic Neo",sans-serif;color:#1d2433;-webkit-print-color-adjust:exact;print-color-adjust:exact}
.page{width:210mm;height:297mm;padding:12mm 14mm 9mm;position:relative;page-break-after:always;overflow:hidden;display:flex;flex-direction:column}
.top{display:flex;justify-content:space-between;align-items:flex-end;border-bottom:2.5px solid #1F3864;padding-bottom:3.5mm}
.brand{font-size:10pt;color:#2F5597;font-weight:700;letter-spacing:.5px}
h1{margin:1.2mm 0 0;font-size:19pt;color:#1F3864;font-weight:800;letter-spacing:-.5px}
.meta{text-align:right;font-size:9pt;color:#5b6577;line-height:1.6}
.who{display:flex;gap:0;margin:3.5mm 0 3.8mm;border:1px solid #d5dced;border-radius:2.5mm;overflow:hidden}
.who div{flex:1;padding:2mm 4mm;font-size:10.5pt;border-right:1px solid #d5dced}
.who div:last-child{border-right:none}
.who span{display:block;font-size:8pt;color:#6b7588;margin-bottom:.5mm}
h2{font-size:11.5pt;color:#1F3864;margin:0 0 1.8mm;display:flex;align-items:center;gap:2mm}
h2:before{content:"";width:1.2mm;height:4.2mm;background:#2F5597;border-radius:.6mm}
.sec{margin-bottom:3.6mm}
.cards{display:flex;gap:4mm}
.card{flex:1;border:1px solid #d5dced;border-radius:2.5mm;padding:2.6mm 4mm;background:#f7f9fd}
.card .sub{font-size:9.5pt;color:#44506a;font-weight:700}
.card .sc{display:flex;align-items:baseline;gap:1.5mm;margin:1mm 0 .6mm}
.card .sc b{font-size:25pt;color:#1F3864;font-weight:800;line-height:1}
.card .sc i{font-style:normal;color:#6b7588;font-size:10pt}
.card .sc .badge{margin-left:auto}
.card .cap{font-size:8.5pt;color:#5b6577}
.track{position:relative;height:2.6mm;background:#e3e8f2;border-radius:1.3mm;margin:2mm 0 1.4mm}
.fill{position:absolute;left:0;top:0;bottom:0;background:#2F5597;border-radius:1.3mm}
.avg{position:absolute;top:-1mm;bottom:-1mm;width:.6mm;background:#1d2433}
.badge{display:inline-block;font-size:8.5pt;font-weight:700;padding:.6mm 2.4mm;border-radius:3mm;color:#fff;white-space:nowrap}
.l0{background:#1F3864}.l1{background:#2F5597}.l2{background:#3f8a6b}.l3{background:#c27b2c}.lx{background:#9aa3b5}
table.units{width:100%;border-collapse:collapse;font-size:9.5pt}
table.units td{padding:1.05mm 1.5mm;border-bottom:1px solid #e6eaf2;vertical-align:middle}
table.units td.s{width:19mm;color:#6b7588;font-size:8.5pt}
table.units td.u{width:33mm;font-weight:600}
table.units td.b{width:auto}
table.units td.p{width:13mm;text-align:right;font-weight:700;color:#1F3864}
table.units td.l{width:21mm;text-align:center}
table.units .track{margin:0}
.legend{font-size:8pt;color:#6b7588;margin-top:1.4mm;display:flex;gap:4mm;align-items:center}
.legend .k{display:inline-block;width:5mm;height:2mm;background:#2F5597;border-radius:1mm;vertical-align:middle;margin-right:1mm}
.legend .t{display:inline-block;width:.6mm;height:3mm;background:#1d2433;vertical-align:middle;margin-right:1mm}
.grid{display:flex;align-items:center;gap:2mm;margin-bottom:1.2mm}
.grid .gl{width:19mm;font-size:8.5pt;color:#44506a;font-weight:700}
.grid .cells{display:grid;grid-template-columns:repeat(20,1fr);flex:1;border:1px solid #d5dced;border-radius:1.5mm;overflow:hidden}
.grid .cells div{text-align:center;font-size:8pt;border-right:1px solid #e6eaf2;padding:.5mm 0;line-height:1.35}
.grid .cells div:last-child{border-right:none}
.grid .cells div.sa{background:#f4f1fa}
.grid .cells b{display:block;font-size:10pt}
.o{color:#2F5597}.x{color:#c0504d}.n{color:#9aa3b5}
.two{display:flex;gap:4mm}
.reco{display:flex;border:1.2px solid #1F3864;border-radius:2.5mm;overflow:hidden;height:24mm}
.reco .t{width:30mm;background:#1F3864;color:#fff;font-weight:700;font-size:11pt;display:flex;align-items:center;justify-content:center}
.reco .w{flex:1}
.op{flex:1;border:1px solid #d5dced;border-radius:2.5mm;padding:3mm 4mm;font-size:9pt;line-height:1.6}
.op ul{margin:1.5mm 0 0;padding-left:4.5mm}
.op li{margin-bottom:.6mm}
.op .tc{margin-top:2mm;padding-top:2mm;border-top:1px dashed #cfd6e4}
.op .tc b{color:#1F3864}
.foot{margin-top:auto;border-top:1px solid #d5dced;padding-top:2.4mm;font-size:8.5pt;color:#5b6577;display:flex;justify-content:space-between;gap:6mm}
.foot b{color:#1F3864}
/* 상담 카드 */
.int{position:absolute;top:8mm;right:14mm;font-size:8pt;color:#c0504d;font-weight:700;border:1px solid #c0504d;padding:.6mm 2mm;border-radius:1mm}
.kpi{display:grid;grid-template-columns:repeat(4,1fr);gap:2.5mm;margin-bottom:3mm}
.kpi div{border:1px solid #d5dced;border-radius:2mm;padding:2mm 3mm;font-size:8.5pt;color:#6b7588}
.kpi b{display:block;font-size:14pt;color:#1d2433;margin-top:.5mm}
.kpi .warn{border-color:#e3a35b;background:#fff6ea}
table.t{width:100%;border-collapse:collapse;font-size:8.8pt}
table.t th{background:#eef2f9;color:#1F3864;font-weight:700;padding:.9mm;border:1px solid #d5dced}
table.t td{padding:.75mm 1.5mm;border:1px solid #e1e6ef;text-align:center}
table.t td.L{text-align:left}
.flag{color:#c0504d;font-weight:700}
table.t.s{font-size:7.6pt} table.t.s td,table.t.s th{padding:.3mm 1.2mm;line-height:1.35}
.wt{font-size:8.5pt;font-weight:700;color:#44506a;margin-bottom:1mm}
.memo{border:1px solid #d5dced;border-radius:2mm;flex:1;min-height:12mm;padding:2mm 3mm;font-size:8.5pt;color:#9aa3b5;
 background:repeating-linear-gradient(#fff 0 7.6mm,#e6eaf2 7.6mm 7.9mm)}
"""


def pct(v):
    return "-" if v is None else "%d%%" % round(v * 100)


def bar(rate, avg, show):
    if rate is None:
        return '<div class="track"></div>'
    a = '<div class="avg" style="left:%.1f%%"></div>' % (avg * 100) if (show and avg is not None) else ""
    return '<div class="track"><div class="fill" style="width:%.1f%%"></div>%s</div>' % (rate * 100, a)


def badge(lv):
    return '<span class="badge %s">%s</span>' % (LVCLS.get(lv, "lx"), lv or "미응시")


def parent_page(s, cfg, exams, stats):
    e = html.escape
    show = cfg["show_avg"]
    subjects = " · ".join(n for _, n in SUBJ)
    cards = []
    for key, nm in SUBJ:
        full = stats[("full", key)]
        if not s["took"][key]:
            cards.append('<div class="card"><div class="sub">%s</div><div class="sc"><b style="font-size:16pt;color:#9aa3b5">미응시</b></div></div>' % nm)
            continue
        sc = s["score"][key]
        n_ok = sum(s["ok"][key])
        avg = stats[("avg", key)]
        lv = level(sc / full if full else 0, cfg)
        cap = "20문항 중 %d문항 정답" % n_ok + (" · 전체 평균 %.1f점" % avg if show and avg is not None else "")
        cards.append('<div class="card"><div class="sub">%s</div><div class="sc"><b>%s</b><i>/ %d점</i>%s</div>%s<div class="cap">%s</div></div>'
                     % (nm, ("%g" % sc), full, badge(lv), bar(sc / full if full else 0, (avg / full) if (avg is not None and full) else None, show), cap))
    rows = []
    for key, nm in SUBJ:
        for u in cfg["units"][key]:
            r = s["unit"][(key, u)]
            rows.append('<tr><td class="s">%s</td><td class="u">%s</td><td class="b">%s</td><td class="p">%s</td><td class="l">%s</td></tr>'
                        % (nm, e(u), bar(r, stats[("uavg", key, u)], show), pct(r) if r is not None else "-", badge(level(r, cfg))))
    legend = '<div class="legend"><span><i class="k"></i>학생 득점률</span>%s<span>성취 수준: 매우 우수 %g%% 이상 · 우수 %g%% 이상 · 양호 %g%% 이상</span></div>' % (
        '<span><i class="t"></i>전체 평균</span>' if show else "", *cfg["lv"])
    grids = []
    for key, nm in SUBJ:
        cells = []
        for q, o in zip(exams[key], s["ok"][key]):
            if not s["took"][key]:
                mark = '<b class="n">-</b>'
            else:
                mark = '<b class="o">○</b>' if o else '<b class="x">×</b>'
            cells.append('<div class="%s">%d%s</div>' % ("sa" if q["no"] > 14 else "", q["no"], mark))
        grids.append('<div class="grid"><div class="gl">%s</div><div class="cells">%s</div></div>' % (nm, "".join(cells)))
    text, tips = opinion(s, cfg, stats)
    tip_html = "".join("<li><b>%s</b> — %s</li>" % (e(u), e(t)) for u, t in tips if t)
    return f"""<div class="page">
<div class="top"><div><div class="brand">{e(cfg['academy'])}</div><h1>{e(cfg['title'])}</h1></div>
<div class="meta">시험일 {e(str(cfg['date']))}<br>평가 영역 {subjects}</div></div>
<div class="who"><div><span>이름</span><b>{e(s['name'])}</b></div><div><span>학교</span>{e(str(s['school']))}</div><div><span>응시 과목</span>{' · '.join(n for k, n in SUBJ if s['took'][k]) or '-'}</div></div>
<div class="sec"><h2>과목별 결과</h2><div class="cards">{''.join(cards)}</div></div>
<div class="sec"><h2>단원별 성취도</h2><table class="units">{''.join(rows)}</table>{legend}</div>
<div class="sec"><h2>문항별 결과</h2>{''.join(grids)}<div class="legend"><span>○ 정답 · × 오답</span><span>1~14번 객관식 · 15~20번 단답형(음영)</span></div></div>
<div class="sec"><h2>종합 의견</h2><div class="op">{e(text)}{('<ul>' + tip_html + '</ul>') if tip_html else ''}</div></div>
<div class="sec"><h2>추천 반</h2><div class="reco"><div class="t">추천 반</div><div class="w"></div></div></div>
<div class="foot"><div>{e(cfg['notice'])}</div><div style="white-space:nowrap"><b>{e(cfg['academy'])}</b> {e(str(cfg['contact']))}</div></div>
</div>"""


def counsel_page(s, cfg, exams, stats):
    e = html.escape
    # 기준 점수와의 차이
    gap = []
    for k, (key, nm) in enumerate(SUBJ):
        sc = s["score"][key]
        if sc is None:
            continue
        ths = sorted({v for r_ in cfg["rules"] for v in (r_[1:3] if k == 0 else r_[3:5]) if v is not None})
        if ths:
            v = min(ths, key=lambda t: abs(sc - t))
            gap.append("%s %g점 (기준 %g점 대비 %+g)" % (nm, sc, v, sc - v))
    near = s["near"]
    manual = s["manual"] and s["manual"] != s["auto"]
    kpi = [("공통수학1", "-" if s["score"]["공수1"] is None else "%g점" % s["score"]["공수1"], False),
           ("공통수학2", "-" if s["score"]["공수2"] is None else "%g점" % s["score"]["공수2"], False),
           ("환산총점 / 석차(참고)", "-" if s["total"] is None else "%g점 · %d/%d" % (s["total"], s["rank"], stats["n"]), False),
           ("자동배정 → 최종반", "%s → %s%s" % (s["auto"] or "-", s["final"] or "-", " (수동)" if manual else ""), near)]
    kp = "".join('<div class="%s">%s<b>%s</b></div>' % ("warn" if w else "", e(a), e(b)) for a, b, w in kpi)
    urows = []
    for key, nm in SUBJ:
        for u in cfg["units"][key]:
            r = s["unit"][(key, u)]
            urows.append("<tr><td>%s</td><td class='L'>%s</td><td><b>%s</b></td><td>%s</td><td>%s</td><td>%s</td></tr>" % (
                nm, e(u), pct(r), pct(stats[("uavg", key, u)]), pct(stats.get(("cavg", key, u, s["final"]))) if s["final"] else "-", level(r, cfg) or "미응시"))
    wrong = {key: [] for key, _ in SUBJ}
    easy_miss = 0
    for key, nm in SUBJ:
        if not s["took"][key]:
            continue
        for q, o in zip(exams[key], s["ok"][key]):
            if o:
                continue
            flag = q["rate"] is not None and q["rate"] >= 0.7
            easy_miss += flag
            wrong[key].append("<tr><td>%d</td><td>%s</td><td class='L'>%s</td><td class='%s'>%s</td></tr>" % (
                q["no"], q["level"], e(POINT.get((key, q["no"]), "") or str(q["unit"])), "flag" if flag else "", pct(q["rate"])))
    wtables = []
    for key, nm in SUBJ:
        body = "".join(wrong[key]) or ("<tr><td colspan=4>%s</td></tr>" % ("미응시" if not s["took"][key] else "틀린 문항 없음"))
        wtables.append("<div style='flex:1'><div class='wt'>%s</div><table class='t s'><tr><th>번호</th><th>난도</th><th>평가 요소</th><th>정답률</th></tr>%s</table></div>" % (nm, body))
    grids = []
    for key, nm in SUBJ:
        cells = []
        for q, o in zip(exams[key], s["ok"][key]):
            mark = '<b class="n">-</b>' if not s["took"][key] else ('<b class="o">○</b>' if o else '<b class="x">×</b>')
            cells.append('<div class="%s">%d%s</div>' % ("sa" if q["no"] > 14 else "", q["no"], mark))
        grids.append('<div class="grid"><div class="gl">%s</div><div class="cells">%s</div></div>' % (nm, "".join(cells)))
    pts = []
    if near:
        pts.append("기준 점수 경계에 있습니다. 단원별 성취와 학습 태도를 함께 보고 반을 확정하세요.")
    if s["auto"] == cfg["classes"][3]:
        pts.insert(0, "자동배정 보류: 과목별 기준에 해당하지 않습니다. 상담 후 반을 정해 엑셀 수동조정 칸에 입력하세요.")
    if easy_miss:
        pts.append("정답률 70%% 이상인 문항에서 %d문항 오답 — 개념 누락 또는 계산 실수 여부를 확인하세요." % easy_miss)
    weak = sorted([(u, s["unit"][(k, u)]) for k, _ in SUBJ for u in cfg["units"][k] if s["unit"][(k, u)] is not None], key=lambda x: x[1])[:2]
    if weak:
        pts.append("상대적 취약 단원: " + ", ".join("%s(%s)" % (u, pct(r)) for u, r in weak))
    if manual:
        pts.append("수동 조정됨: 자동배정 %s → %s" % (s["auto"], s["final"]))
    if s["memo"]:
        pts.append("비고: %s" % s["memo"])
    return f"""<div class="page"><div class="int">내부용 · 외부 유출 금지</div>
<div class="top"><div><div class="brand">{e(cfg['academy'])} · 개별 상담 카드</div><h1>{e(s['name'])} <span style="font-size:11pt;color:#5b6577;font-weight:500">{e(str(s['school']))} · {e(str(s['kind']))} · {e(str(s['phone']))}</span></h1></div>
<div class="meta">{e(cfg['title'])}<br>{' / '.join(gap) if gap else ''}</div></div>
<div style="height:4mm"></div><div class="kpi">{kp}</div>
<div class="sec"><h2>상담 포인트</h2><ul style="margin:0;padding-left:5mm;font-size:9.3pt;line-height:1.7">{''.join('<li>%s</li>' % e(p) for p in pts) or '<li>특이 사항 없음</li>'}</ul></div>
<div class="sec"><h2>단원별 득점률</h2><table class="t"><tr><th>과목</th><th>단원</th><th>학생</th><th>전체 평균</th><th>최종반 평균</th><th>수준</th></tr>{''.join(urows)}</table></div>
<div class="sec"><h2>문항별 정오</h2>{''.join(grids)}</div>
<div class="sec"><h2>틀린 문항 <span style="font-size:8.5pt;color:#6b7588;font-weight:400">(빨간 정답률 = 70% 이상이 맞힌 문항)</span></h2>
<div style="display:flex;gap:4mm">{''.join(wtables)}</div></div>
<div class="sec" style="flex:1;display:flex;flex-direction:column;margin-bottom:0"><h2>상담 메모</h2><div class="memo"></div></div>
</div>"""


def render(pages, out_pdf):
    from playwright.sync_api import sync_playwright
    doc = '<!doctype html><html><head><meta charset="utf-8"><style>%s</style></head><body>%s</body></html>' % (CSS, "".join(pages))
    tmp = os.path.abspath(out_pdf) + ".html"
    open(tmp, "w", encoding="utf-8").write(doc)
    with sync_playwright() as p:
        kw = {"executable_path": CHROME} if os.path.exists(CHROME) else {}
        b = p.chromium.launch(args=["--disable-background-networking", "--no-pings"], **kw)
        pg = b.new_page()
        pg.goto("file://" + tmp)
        pg.pdf(path=out_pdf, format="A4", print_background=True, margin={"top": "0", "bottom": "0", "left": "0", "right": "0"})
        b.close()
    os.remove(tmp)


def main(xlsx, outdir):
    cfg, exams, students = read(xlsx)
    stats = compute(cfg, exams, students)
    os.makedirs(os.path.join(outdir, "학부모보고서"), exist_ok=True)
    targets = [s for s in students if s["total"] is not None]
    pages = [parent_page(s, cfg, exams, stats) for s in targets]
    render(pages, os.path.join(outdir, "학부모보고서_전체.pdf"))
    for s, pgx in zip(targets, pages):
        safe = re.sub(r'[\\/:*?"<>|]', "_", s["name"])
        render([pgx], os.path.join(outdir, "학부모보고서", "%02d_%s.pdf" % (s["no"], safe)))
    render([counsel_page(s, cfg, exams, stats) for s in targets], os.path.join(outdir, "상담카드_전체.pdf"))
    # 수정 가능한 한글(HML) 보고서: 전체 1개 + 학생별
    import hml_report
    hdir = os.path.join(outdir, "학부모보고서_한글")
    os.makedirs(hdir, exist_ok=True)
    hml_report.build(targets, cfg, exams, stats, SUBJ, level, opinion, os.path.join(hdir, "전체.hml"))
    for s in targets:
        safe = re.sub(r'[\\/:*?"<>|]', "_", s["name"])
        hml_report.build([s], cfg, exams, stats, SUBJ, level, opinion, os.path.join(hdir, "%02d_%s.hml" % (s["no"], safe)))
    print("학생 %d명: 학부모 보고서·상담 카드 생성 완료 → %s" % (len(targets), outdir))


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
