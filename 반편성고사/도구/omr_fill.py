# -*- coding: utf-8 -*-
"""답안지 판독 결과(omr_read.py)를 채점 엑셀에 넣는다 (객관식 1~14번만, 단답형 15~20번은 직접 1/0 입력)
사용: python omr_fill.py <채점엑셀.xlsx> <공수1|공수2> <판독결과.json> <짝맞춤.csv> [수정.csv]
- 짝맞춤.csv : 쪽,학생   (쪽 = 판독결과의 번호, 학생 = 학생명단의 No 또는 이름)
- 수정.csv   : 쪽,문항,답 (확인 그림을 보고 고친 답. 답을 비우면 표시 없음, 0이면 오답 처리.
               15~20번은 채점 결과 1/0을 적으면 함께 입력)
- 표시 없음은 빈칸, 두 개 이상 표시는 0(오답)으로 넣는다. 넣기 전 원본은 *_백업.xlsx 로 저장
"""
import sys, os, csv, json, shutil
from openpyxl import load_workbook

FIRST, LAST, N_MC = 7, 86, 14

def read_csv(path):
    with open(path, encoding="utf-8-sig") as f:
        return [r for r in csv.reader(f) if r and not r[0].strip().startswith("#") and r[0].strip() != "쪽"]

def main(xlsx, subj, result, pairing, fixes=None):
    assert subj in ("공수1", "공수2"), "과목은 공수1 또는 공수2"
    res = {r["no"]: r for r in json.load(open(result, encoding="utf-8"))}
    wb = load_workbook(xlsx)
    roster = wb["학생명단"]
    by_name, by_no = {}, {}
    for r in range(FIRST, LAST + 1):
        nm = roster.cell(r, 2).value
        no = r - FIRST + 1
        by_no[no] = r
        if nm:
            by_name.setdefault(str(nm).strip(), []).append(r)
    fix = {}
    for row in read_csv(fixes) if fixes else []:
        page, q, v = int(row[0]), int(row[1]), row[2].strip() if len(row) > 2 else ""
        fix[(page, q)] = int(v) if v else None
    ws = wb[subj]
    done, problems = [], []
    for row in read_csv(pairing):
        page, who = int(row[0]), row[1].strip()
        if page not in res or "error" in res[page]:
            problems.append("%d쪽: 판독 결과 없음" % page); continue
        if who.isdigit():
            r = by_no.get(int(who))
        else:
            rows = by_name.get(who, [])
            if len(rows) != 1:
                problems.append("%d쪽: 학생명단에서 '%s' %s" % (page, who, "없음" if not rows else "동명이인 — No로 지정")); continue
            r = rows[0]
        item = res[page]
        for q in range(N_MC):
            a, fl = item["answers"][q], item["flags"][q]
            v = 0 if fl.startswith("두 개") else a
            if (page, q + 1) in fix:
                v = fix[(page, q + 1)]
            ws.cell(r, 4 + q).value = v
        for q in range(N_MC + 1, 21):          # 단답형 15~20번: 수정.csv에 1/0을 적은 경우만
            if (page, q) in fix:
                ws.cell(r, 3 + q).value = fix[(page, q)]
        done.append((page, roster.cell(r, 2).value))
    shutil.copy(xlsx, xlsx.replace(".xlsx", "_백업.xlsx"))
    wb.save(xlsx)
    print("%s: %d명 입력 완료" % (subj, len(done)))
    for p in problems:
        print("  확인:", p)

if __name__ == "__main__":
    main(*sys.argv[1:])
