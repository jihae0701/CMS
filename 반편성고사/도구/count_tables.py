# -*- coding: utf-8 -*-
"""반배정 시트 오른쪽(S열~)의 '요일반별 인원(자사 구분)'·'선행과목별 인원' 표
build_workbook.py 가 새 양식을 만들 때 쓰고, 이미 쓰고 있는 채점 엑셀에 표만 다시 넣을 때도 씀:
    python count_tables.py <채점엑셀.xlsx>
"""
import sys
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L

FONT = "맑은 고딕"
FIRST, LAST = 7, 156
ROW0, COL0 = 20, 19          # S20부터
PRE_DEFAULT = ["미적2", "미적1", "공수1심화", None, None]   # 선행과목 이름(노란 칸, 바꿔 쓸 수 있음)

def f(bold=False, size=10, color="000000"):
    return Font(name=FONT, bold=bold, size=size, color=color)

INPUT = PatternFill("solid", fgColor="FFF2CC")
HEAD = PatternFill("solid", fgColor="1F3864")
GREY = PatternFill("solid", fgColor="F2F2F2")
CLS4 = ["DDEBF7", "E2EFDA", "FCE4D6", "E7E6E6"]
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
C = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

def cell(ws, ref, v=None, font=None, fill=None, align=C, border=True, fmt=None):
    c = ws[ref]
    c.value = v
    c.font = font or f()
    c.fill = fill or PatternFill()
    c.alignment = align
    c.border = BOX if border else Border()
    if fmt:
        c.number_format = fmt
    return c

def header(ws, row, labels, start):
    for i, lab in enumerate(labels):
        cell(ws, f"{L(start + i)}{row}", lab, f(True, 10, "FFFFFF"), HEAD)

def add_count_tables(b, first=FIRST, last=LAST):
    O = f"$O${first}:$O${last}"       # 공통수학반(최종)
    N = f"$N${first}:$N${last}"       # 선행반(입력)
    B = f"$B${first}:$B${last}"       # 이름
    G = f"$G${first}:$G${last}"       # 공수1
    H = f"$H${first}:$H${last}"       # 공수2
    cs = [L(COL0 + i) for i in range(7)]
    for rng in list(b.merged_cells.ranges):          # 예전 표 자리 비우기
        if rng.min_row >= ROW0 and rng.min_col >= COL0:
            b.unmerge_cells(str(rng))
    for r in range(ROW0, ROW0 + 45):
        for c in cs:
            cell(b, f"{c}{r}", None, border=False)

    # ① 요일반별 인원: 일반 / 자사(반 이름 뒤 괄호, 예 TF1월금(자사)) 구분
    r0 = ROW0
    cell(b, f"{cs[0]}{r0}", "요일반별 인원 (최종 기준 · 자사 구분)", f(True, 11, "1F3864"), align=LEFT, border=False)
    header(b, r0 + 1, ["반", "일반", "자사", "합계", "공수1\n평균", "공수2\n평균"], COL0)
    pfx = lambda k: f"SUBSTITUTE('설정'!$B${6 + k},\"반\",\"\")"
    rows = {}
    r = r0 + 2
    for k, suf in ((0, "월금"), (0, "화목"), (2, "월금"), (2, "화목")):
        a, t, u, v, w, x = [f"{c}{r}" for c in cs[:6]]
        cell(b, a, f"={pfx(k)}&\"{suf}\"", f(True), PatternFill("solid", fgColor=CLS4[k]))
        cell(b, t, f"=COUNTIF({O},{a})")
        cell(b, u, f"=COUNTIF({O},{a}&\"(*\")")
        cell(b, v, f"={t}+{u}", f(True))
        for col, rg in ((w, G), (x, H)):
            cell(b, col, f"=IFERROR((SUMIF({O},{a},{rg})+SUMIF({O},{a}&\"(*\",{rg}))/(COUNTIFS({O},{a},{rg},\">=0\")"
                         f"+COUNTIFS({O},{a}&\"(*\",{rg},\">=0\")),\"\")", fmt="0.0")
        rows[(k, suf)] = r
        r += 1
    for suf in ("월금", "화목"):                       # TF1+TF2 요일 소계
        r1, r2 = rows[(0, suf)], rows[(2, suf)]
        cell(b, f"{cs[0]}{r}", f"TF {suf} 소계", f(True), GREY)
        for c in cs[1:4]:
            cell(b, f"{c}{r}", f"={c}{r1}+{c}{r2}", f(True), GREY)
        for c in cs[4:6]:
            cell(b, f"{c}{r}", None, fill=GREY)
        rows[suf] = r
        r += 1
    a = f"{cs[0]}{r}"                                  # HF반
    cell(b, a, f"={pfx(1)}", f(True), PatternFill("solid", fgColor=CLS4[1]))
    cell(b, f"{cs[2]}{r}", f"=COUNTIF({O},{a}&\"*(*\")")
    cell(b, f"{cs[3]}{r}", f"=COUNTIF({O},{a}&\"*\")", f(True))
    cell(b, f"{cs[1]}{r}", f"={cs[3]}{r}-{cs[2]}{r}")
    cell(b, f"{cs[4]}{r}", f"=IFERROR(AVERAGEIF({O},{a}&\"*\",{G}),\"\")", fmt="0.0")
    cell(b, f"{cs[5]}{r}", f"=IFERROR(AVERAGEIF({O},{a}&\"*\",{H}),\"\")", fmt="0.0")
    rows["HF"] = r
    r += 1
    singles = [("=IF('설정'!$B$10=\"\",\"학교 지정\",'설정'!$B$10)&\"(요일 미정)\"", "'설정'!$B$10&\"(*\"", True),
               ("보류", "\"보류*\"", False), ("퇴원", "\"퇴원*\"", False)]
    for lab, crit, school in singles:
        cell(b, f"{cs[0]}{r}", lab, f(True))
        cell(b, f"{cs[1]}{r}", None if school else f"=COUNTIF({O},{crit})")
        cell(b, f"{cs[2]}{r}", f"=COUNTIF({O},{crit})" if school else None)
        cell(b, f"{cs[3]}{r}", f"=COUNTIF({O},{crit})", f(True))
        cell(b, f"{cs[4]}{r}", None)
        cell(b, f"{cs[5]}{r}", None)
        rows.setdefault("singles", []).append(r)
        r += 1
    tot = r + 1
    known = "+".join(f"{cs[3]}{x}" for x in [rows["월금"], rows["화목"], rows["HF"]] + rows["singles"])
    cell(b, f"{cs[0]}{r}", "기타(확인 필요)", f(True, color="C00000"))
    for c in (cs[1], cs[2], cs[4], cs[5]):
        cell(b, f"{c}{r}", None)
    cell(b, f"{cs[3]}{r}", f"={cs[3]}{tot}-({known})", f(True, color="C00000"))
    cell(b, f"{cs[0]}{tot}", "합계", f(True), GREY)
    cell(b, f"{cs[1]}{tot}", f"=SUM({cs[1]}{r0 + 2}:{cs[1]}{r0 + 5})+{cs[1]}{rows['HF']}+" + "+".join(f"{cs[1]}{x}" for x in rows["singles"]), f(True), GREY)
    cell(b, f"{cs[2]}{tot}", f"=SUM({cs[2]}{r0 + 2}:{cs[2]}{r0 + 5})+{cs[2]}{rows['HF']}+" + "+".join(f"{cs[2]}{x}" for x in rows["singles"]), f(True), GREY)
    cell(b, f"{cs[3]}{tot}", f"=COUNTIF({O},\"?*\")", f(True), GREY)
    cell(b, f"{cs[4]}{tot}", f"=IFERROR(AVERAGE({G}),\"\")", f(True), GREY, fmt="0.0")
    cell(b, f"{cs[5]}{tot}", f"=IFERROR(AVERAGE({H}),\"\")", f(True), GREY, fmt="0.0")
    note = tot + 1
    cell(b, f"{cs[0]}{note}", "· 자사 = 반 이름 뒤에 괄호가 붙은 학생(예: TF1월금(자사))\n"
         "· 학교 지정(요일 미정) = 공통수학반(입력)이 비어 'TF자사(TF1)'처럼 자동으로 들어간 학생\n"
         "· 기타 = 위 반 이름과 맞지 않는 입력(오타 등) → 공통수학반(입력) 확인", f(size=9, color="595959"), align=LEFT, border=False)
    b.merge_cells(f"{cs[0]}{note}:{cs[6]}{note + 2}")

    # ② 선행과목별 인원 (선행반 입력 기준, 공통수학반 요일별)
    p0 = note + 4
    cell(b, f"{cs[0]}{p0}", "선행과목별 인원 (선행반 입력 기준)", f(True, 11, "1F3864"), align=LEFT, border=False)
    header(b, p0 + 1, ["선행과목", "인원", "TF\n월금", "TF\n화목", "HF", "기타"], COL0)
    r = p0 + 2
    for nm in PRE_DEFAULT:
        a, t, u, v, w, x = [f"{c}{r}" for c in cs[:6]]
        cell(b, a, nm, f(True), INPUT)
        cell(b, t, f"=IF({a}=\"\",\"\",COUNTIF({N},{a}))", f(True))
        cell(b, u, f"=IF({a}=\"\",\"\",COUNTIFS({N},{a},{O},\"TF*월금*\"))")
        cell(b, v, f"=IF({a}=\"\",\"\",COUNTIFS({N},{a},{O},\"TF*화목*\"))")
        cell(b, w, f"=IF({a}=\"\",\"\",COUNTIFS({N},{a},{O},\"HF*\"))")
        cell(b, x, f"=IF({a}=\"\",\"\",{t}-{u}-{v}-{w})")
        r += 1
    last_in = r - 1
    a = f"{cs[0]}{r}"
    cell(b, a, "목록에 없는 과목", f(True, color="C00000"))
    cell(b, f"{cs[1]}{r}", f"=COUNTIF({N},\"?*\")-SUM({cs[1]}{p0 + 2}:{cs[1]}{last_in})", f(True, color="C00000"))
    for c in cs[2:6]:
        cell(b, f"{c}{r}", None)
    r += 1
    cell(b, f"{cs[0]}{r}", "미입력", f(True))
    cell(b, f"{cs[1]}{r}", f"=COUNTIFS({B},\"?*\",{N},\"\")", f(True))
    cell(b, f"{cs[2]}{r}", f"=COUNTIFS({B},\"?*\",{N},\"\",{O},\"TF*월금*\")")
    cell(b, f"{cs[3]}{r}", f"=COUNTIFS({B},\"?*\",{N},\"\",{O},\"TF*화목*\")")
    cell(b, f"{cs[4]}{r}", f"=COUNTIFS({B},\"?*\",{N},\"\",{O},\"HF*\")")
    cell(b, f"{cs[5]}{r}", f"={cs[1]}{r}-{cs[2]}{r}-{cs[3]}{r}-{cs[4]}{r}")
    r += 1
    cell(b, f"{cs[0]}{r}", "합계", f(True), GREY)
    for c in cs[1:6]:
        cell(b, f"{c}{r}", f"=SUM({c}{p0 + 2}:{c}{r - 1})", f(True), GREY)
    cell(b, f"{cs[0]}{r + 1}", "· 선행과목 이름(노란 칸)은 바꾸거나 더 적을 수 있음 · TF월금/TF화목은 자사 학생 포함"
         " · 기타 = 학교 지정(요일 미정)·보류·퇴원 등", f(size=9, color="595959"), align=LEFT, border=False)
    b.merge_cells(f"{cs[0]}{r + 1}:{cs[6]}{r + 2}")
    b.column_dimensions[cs[0]].width = 16


if __name__ == "__main__":
    path = sys.argv[1]
    wb = load_workbook(path)
    add_count_tables(wb["반배정"])
    wb.save(path)
    print("반배정 인원표 갱신 →", path)
